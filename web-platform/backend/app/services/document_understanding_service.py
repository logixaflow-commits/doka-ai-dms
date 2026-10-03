from __future__ import annotations

import json
import mimetypes
import re
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict

from loguru import logger

from app.services.metadata_extractor import metadata_extractor
from app.services.safe_workspace_service import SafeWorkspaceService, safe_workspace_service, sha256_file


TEXT_EXTENSIONS = {
    ".txt", ".csv", ".tsv", ".json", ".xml", ".md", ".log", ".yaml", ".yml",
    ".html", ".htm", ".py", ".js", ".ts", ".tsx", ".jsx", ".sql",
}


class DocumentUnderstandingService:
    """Local-only content extraction and metadata enrichment. AI is not used here."""

    def __init__(self, workspace: SafeWorkspaceService):
        self.workspace = workspace

    def _extract_text(self, path: Path) -> tuple[str, str]:
        ext = path.suffix.lower()
        if ext in TEXT_EXTENSIONS:
            for encoding in ("utf-8", "utf-16", "cp1252"):
                try:
                    return path.read_text(encoding=encoding, errors="strict"), "text"
                except (UnicodeDecodeError, UnicodeError):
                    continue
            return "", "text_unreadable"

        if ext == ".docx":
            try:
                with zipfile.ZipFile(path) as archive:
                    xml = archive.read("word/document.xml")
                root = ET.fromstring(xml)
                text = " ".join(node.text or "" for node in root.iter() if node.tag.endswith("}t"))
                return re.sub(r"\s+", " ", text).strip(), "docx"
            except Exception as exc:
                logger.warning("DOCX extraction failed for %s: %s", path, exc)
                return "", "docx_failed"

        if ext in {".xlsx", ".xlsm"}:
            try:
                from openpyxl import load_workbook
                workbook = load_workbook(path, read_only=True, data_only=True)
                chunks = []
                for sheet in workbook.worksheets:
                    chunks.append(f"[Sheet: {sheet.title}]")
                    for row in sheet.iter_rows(values_only=True):
                        values = [str(value) for value in row if value is not None]
                        if values:
                            chunks.append(" | ".join(values))
                workbook.close()
                return "\n".join(chunks), "xlsx"
            except Exception as exc:
                logger.warning("Excel extraction failed for %s: %s", path, exc)
                return "", "xlsx_failed"

        if ext in {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}:
            try:
                from app.services.ocr_service import ocr_service
                mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
                text, language = ocr_service.process_file(str(path), mime)
                return text or "", language or "unknown"
            except Exception as exc:
                logger.warning("OCR failed for %s: %s", path, exc)
                return "", "ocr_failed"

        return "", "not_text_extracted"

    def analyze(self, session_id: str) -> Dict[str, Any]:
        with self.workspace._lock(session_id):
            return self._analyze_locked(session_id)

    def _analyze_locked(self, session_id: str) -> Dict[str, Any]:
        session = self.workspace._dir(session_id)
        manifest = self.workspace._read(session / "manifest.json")
        if not manifest:
            raise ValueError(f"Unknown import session: {session_id}")

        _source, root = self.workspace.validate_manifest_paths(session_id, manifest)
        if not root.is_dir():
            raise ValueError("Working copy does not exist.")

        results = []
        for relative_path, item in manifest.get("files", {}).items():
            if not item.get("verified"):
                continue
            if item.get("relative_path") != relative_path:
                raise ValueError("Import manifest contains an inconsistent relative path.")
            path = (root / Path(relative_path)).resolve()
            try:
                path.relative_to(root)
            except ValueError as exc:
                raise ValueError("Import manifest file path is outside the working copy.") from exc
            if not path.is_file():
                continue
            expected_hash = str(item.get("sha256", "")).lower()
            if len(expected_hash) != 64 or sha256_file(path).lower() != expected_hash:
                raise ValueError(f"Working-copy integrity check failed: {relative_path}")
            text, method = self._extract_text(path)
            metadata: Dict[str, Any] = {}
            if text.strip():
                try:
                    metadata = metadata_extractor.extract(text)
                except Exception as exc:
                    logger.warning("Metadata extraction failed for %s: %s", path, exc)
            results.append({
                "relative_path": item["relative_path"],
                "filename": item["filename"],
                "extension": item["extension"],
                "sha256": item["sha256"],
                "extraction_method": method,
                "text_length": len(text),
                "language": method if method in {"mya", "eng", "mya+eng"} else None,
                "metadata": metadata,
                "text_preview": text[:2000],
            })

        result = {
            "schema_version": 1,
            "session_id": session_id,
            "analyzed_files": len(results),
            "analyzed_at": datetime.now(timezone.utc).isoformat(),
            "ai_used": False,
            "results": results,
        }
        self.workspace._write(session / "understanding.json", result)
        return {k: v for k, v in result.items() if k != "results"}


document_understanding_service = DocumentUnderstandingService(safe_workspace_service)
